"""
CRM — Contacts enrichis
Couvre : liste contacts, fiche unifiée, notes, valeur client, inactivité, import/export CSV et Excel.
Tags et reminders sont dans tags.py et reminders.py (routeurs séparés, même préfixe /contacts).
"""
import csv
import io
from datetime import datetime, timedelta, timezone
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import Response
from pydantic import BaseModel, field_validator

from app.core.supabase import get_supabase_admin
from app.middleware.tenant import get_current_tenant, get_current_user
from app.api.v1.attachments import PHOTO_BUCKET, _signed_url_bucket
from app.services.activity import log_activity
from app.services.phone import clean_phone, is_valid_phone, validate_phone_field

router = APIRouter(prefix="/contacts", tags=["Contacts"])

INACTIVE_DAYS = 180  # 6 mois sans RDV confirmé → inactif


# ── Schémas ──────────────────────────────────────────────────────────────────

class ContactUpdate(BaseModel):
    notes: Optional[str] = None
    first_name: Optional[str] = None
    last_name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    category: Optional[str] = None
    segment: Optional[str] = None
    contact_details: Optional[dict] = None

    @field_validator("phone")
    @classmethod
    def validate_phone(cls, v: Optional[str]) -> Optional[str]:
        return validate_phone_field(v)


CATEGORIES = {"client", "prospect", "partenaire", "fournisseur", "autre"}


class ActivityCreate(BaseModel):
    type: str = "note"
    content: str


# ── CSV export/import — colonnes alignées sur la fiche contact (contact_details JSONB) ───────

CSV_COLUMNS = [
    "first_name", "last_name", "first_name_2", "first_name_3",
    "gender", "title", "nickname",
    "email", "phone", "email_pro", "email_perso", "phone_pro", "phone_perso", "phone_preferred",
    "category", "segment",
    "country_residence", "country_origin", "birthday_day", "birthday_month",
    "street_number", "street", "city", "postal_code", "region",
    "website", "linkedin", "instagram", "facebook",
    "notes", "source", "created_at",
]

# Alias tolérants acceptés en import (colonne CSV → clé canonique), en plus du nom exact.
CSV_COLUMN_ALIASES = {
    "prenom": "first_name", "prénom": "first_name",
    "nom": "last_name",
    "telephone": "phone", "téléphone": "phone",
    "email_pro": "email_pro", "email_perso": "email_perso",
    "tel_pro": "phone_pro", "telephone_pro": "phone_pro", "téléphone_pro": "phone_pro",
    "tel_perso": "phone_perso", "telephone_perso": "phone_perso", "téléphone_perso": "phone_perso",
    "categorie": "category", "catégorie": "category",
    "pays_residence": "country_residence", "pays_de_residence": "country_residence",
    "pays_origine": "country_origin", "pays_d_origine": "country_origin",
    "jour_naissance": "birthday_day", "mois_naissance": "birthday_month",
    "numero_rue": "street_number", "n_rue": "street_number",
    "rue": "street", "ville": "city", "code_postal": "postal_code", "region": "region",
    "site_internet": "website", "site_web": "website",
}


def _contact_to_csv_row(c: dict) -> list:
    d = c.get("contact_details") or {}
    address = d.get("address") or {}
    birthday = d.get("birthday") or {}
    urls = d.get("urls") or {}
    values = {
        "first_name": c.get("first_name"), "last_name": c.get("last_name"),
        "first_name_2": d.get("first_name_2"), "first_name_3": d.get("first_name_3"),
        "gender": d.get("gender"), "title": d.get("title"), "nickname": d.get("nickname"),
        "email": c.get("email"), "phone": c.get("phone"),
        "email_pro": d.get("email_pro"), "email_perso": d.get("email_perso"),
        "phone_pro": d.get("phone_pro"), "phone_perso": d.get("phone_perso"),
        "phone_preferred": d.get("phone_preferred"),
        "category": c.get("category"), "segment": c.get("segment"),
        "country_residence": d.get("country_residence"), "country_origin": d.get("country_origin"),
        "birthday_day": birthday.get("day"), "birthday_month": birthday.get("month"),
        "street_number": address.get("street_number"), "street": address.get("street"),
        "city": address.get("city"), "postal_code": address.get("postal_code"), "region": address.get("region"),
        "website": urls.get("website"), "linkedin": urls.get("linkedin"),
        "instagram": urls.get("instagram"), "facebook": urls.get("facebook"),
        "notes": c.get("notes"), "source": c.get("source"),
        "created_at": (c.get("created_at") or "")[:10],
    }
    return [values.get(col) if values.get(col) is not None else "" for col in CSV_COLUMNS]


def _csv_row_to_contact_fields(row: dict) -> dict:
    """row : dict CSV normalisé (clés en minuscules, sans espaces). Retourne les champs prêts pour l'insert contact."""
    normalized = {}
    for key, value in row.items():
        canonical = CSV_COLUMN_ALIASES.get(key, key)
        if canonical in CSV_COLUMNS and value:
            normalized[canonical] = value

    # Téléphones : on tolère les séparateurs courants (espaces, tirets) dans le CSV et on les nettoie ;
    # si le résultat contient encore des lettres, on abandonne juste ce champ plutôt que de bloquer tout l'import.
    for phone_key in ("phone", "phone_pro", "phone_perso"):
        if phone_key in normalized:
            cleaned = clean_phone(normalized[phone_key])
            normalized[phone_key] = cleaned if is_valid_phone(cleaned) else ""

    category = normalized.get("category", "").lower()
    if category not in CATEGORIES:
        category = None

    contact_details = {
        "first_name_2": normalized.get("first_name_2", ""), "first_name_3": normalized.get("first_name_3", ""),
        "gender": normalized.get("gender", ""), "title": normalized.get("title", ""),
        "nickname": normalized.get("nickname", ""),
        "email_pro": normalized.get("email_pro", ""), "email_perso": normalized.get("email_perso", ""),
        "phone_pro": normalized.get("phone_pro", ""), "phone_perso": normalized.get("phone_perso", ""),
        "phone_preferred": normalized.get("phone_preferred", ""),
        "country_residence": normalized.get("country_residence", ""), "country_origin": normalized.get("country_origin", ""),
        "birthday": {"day": normalized.get("birthday_day") or None, "month": normalized.get("birthday_month") or None},
        "address": {
            "street_number": normalized.get("street_number", ""), "street": normalized.get("street", ""),
            "city": normalized.get("city", ""), "postal_code": normalized.get("postal_code", ""),
            "region": normalized.get("region", ""),
        },
        "urls": {
            "website": normalized.get("website", ""), "linkedin": normalized.get("linkedin", ""),
            "instagram": normalized.get("instagram", ""), "facebook": normalized.get("facebook", ""),
        },
    }
    # Ne garde contact_details que s'il contient au moins une valeur non vide (évite d'écraser avec du vide)
    has_details = any(
        v for v in (*{k: v for k, v in contact_details.items() if k not in ("birthday", "address", "urls")}.values(),
                     *contact_details["birthday"].values(), *contact_details["address"].values(), *contact_details["urls"].values())
    )

    return {
        "first_name": normalized.get("first_name") or None,
        "last_name": normalized.get("last_name") or None,
        "email": (normalized.get("email") or "").lower() or None,
        "phone": normalized.get("phone") or None,
        "notes": normalized.get("notes") or None,
        "category": category,
        "segment": normalized.get("segment") or None,
        "contact_details": contact_details if has_details else None,
    }


def _rows_to_xlsx_bytes(header: list[str], rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(header)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _read_xlsx_rows(content: bytes) -> list[dict]:
    """Lit la première feuille d'un fichier Excel — 1ère ligne = en-têtes, retourne des dicts normalisés (clés en minuscules)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        return []

    rows = []
    for raw_row in rows_iter:
        row = {}
        for key, value in zip(header, raw_row):
            if not key or value is None:
                continue
            row[key] = str(value).strip()
        if row:
            rows.append(row)
    return rows


# ── Helpers ───────────────────────────────────────────────────────────────────

def _enrich_contacts(sb, tenant_id: str, contacts: list[dict]) -> list[dict]:
    """Ajoute tags, valeur client et statut inactivité à chaque contact."""
    if not contacts:
        return []

    contact_ids = [c["id"] for c in contacts]

    # Tags
    tag_links = (
        sb.table("contact_tag_link")
        .select("contact_id, contact_tag(id, name, color)")
        .in_("contact_id", contact_ids)
        .execute()
    ).data or []
    tags_by_contact: dict[str, list] = {}
    for link in tag_links:
        cid = link["contact_id"]
        tag = link.get("contact_tag") or {}
        if tag:
            tags_by_contact.setdefault(cid, []).append(tag)

    # Derniers RDV confirmés + valeur (prix prestation si dispo)
    # appointment n'a pas de tenant_id — sécurité assurée par le filtre contact_id (contacts déjà filtrés par tenant)
    appts = (
        sb.table("appointment")
        .select("contact_id, scheduled_at, service_offer(price_eur)")
        .eq("status", "confirmed")
        .in_("contact_id", contact_ids)
        .order("scheduled_at", desc=True)
        .execute()
    ).data or []

    last_appt: dict[str, str] = {}
    value_by_contact: dict[str, float] = {}
    for a in appts:
        cid = a["contact_id"]
        if cid not in last_appt:
            last_appt[cid] = a["scheduled_at"]
        price = (a.get("service_offer") or {}).get("price_eur") or 0
        value_by_contact[cid] = value_by_contact.get(cid, 0) + float(price)

    cutoff = (datetime.now(timezone.utc) - timedelta(days=INACTIVE_DAYS)).isoformat()

    for c in contacts:
        cid = c["id"]
        c["tags"] = tags_by_contact.get(cid, [])
        c["client_value"] = round(value_by_contact.get(cid, 0), 2)
        last = last_appt.get(cid)
        c["last_appointment_at"] = last
        c["is_inactive"] = (last is None or last < cutoff)

    return contacts


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.get("/export")
async def export_contacts_csv(
    format: str = Query(default="csv", pattern="^(csv|xlsx)$"),
    tenant_id: str = Depends(get_current_tenant),
):
    """Exporte tous les contacts du tenant en CSV ou Excel (téléchargement)."""
    sb = get_supabase_admin()

    contacts = (
        sb.table("contact")
        .select("first_name, last_name, email, phone, category, segment, contact_details, notes, source, created_at")
        .eq("tenant_id", tenant_id)
        .is_("deleted_at", "null")
        .order("created_at", desc=True)
        .execute()
    ).data or []

    rows = [_contact_to_csv_row(c) for c in contacts]

    if format == "xlsx":
        return Response(
            content=_rows_to_xlsx_bytes(CSV_COLUMNS, rows),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=contacts.xlsx"},
        )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(CSV_COLUMNS)
    for row in rows:
        writer.writerow(row)

    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=contacts.csv"},
    )


IMPORT_TEMPLATE_EXAMPLE = [
    "Marie", "Dupont", "", "", "Femme", "", "",
    "marie.dupont@exemple.com", "+33612345678", "", "", "", "", "",
    "client", "", "BE", "", "", "",
    "12", "Rue de la Loi", "Bruxelles", "1000", "",
    "", "", "", "",
    "Cliente fidèle",
]


@router.get("/import-template")
async def download_import_template(
    format: str = Query(default="csv", pattern="^(csv|xlsx)$"),
    tenant_id: str = Depends(get_current_tenant),
):
    """Modèle vide (avec une ligne d'exemple) à remplir pour l'import — colonnes alignées sur CSV_COLUMNS (sans source/created_at, réservés à l'export)."""
    import_columns = [c for c in CSV_COLUMNS if c not in ("source", "created_at")]

    if format == "xlsx":
        return Response(
            content=_rows_to_xlsx_bytes(import_columns, [IMPORT_TEMPLATE_EXAMPLE]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=modele_contacts.xlsx"},
        )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(import_columns)
    writer.writerow(IMPORT_TEMPLATE_EXAMPLE)

    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=modele_contacts.csv"},
    )


@router.get("/")
async def list_contacts(
    q: Optional[str] = None,
    tag_id: Optional[str] = None,
    inactive_only: bool = False,
    category: Optional[str] = None,
    segment: Optional[str] = None,
    limit: int = Query(default=30, le=100),
    offset: int = 0,
    tenant_id: str = Depends(get_current_tenant),
):
    """Liste paginée des contacts avec enrichissement (tags, valeur, inactivité)."""
    sb = get_supabase_admin()

    query = (
        sb.table("contact")
        .select("id, first_name, last_name, email, phone, contact_type, category, segment, created_at")
        .eq("tenant_id", tenant_id)
        .is_("deleted_at", "null")
        .order("created_at", desc=True)
        .range(offset, offset + limit - 1)
    )

    if q:
        query = query.or_(
            f"first_name.ilike.%{q}%,last_name.ilike.%{q}%,email.ilike.%{q}%,phone.ilike.%{q}%"
        )
    if category:
        query = query.eq("category", category)
    if segment:
        query = query.eq("segment", segment)

    contacts = (query.execute()).data or []

    # Filtre tag côté applicatif (supabase-py ne supporte pas le join avec filtre many-to-many facilement)
    if tag_id:
        linked = {
            row["contact_id"]
            for row in (
                sb.table("contact_tag_link")
                .select("contact_id")
                .eq("tag_id", tag_id)
                .execute()
            ).data or []
        }
        contacts = [c for c in contacts if c["id"] in linked]

    contacts = _enrich_contacts(sb, tenant_id, contacts)

    if inactive_only:
        contacts = [c for c in contacts if c["is_inactive"]]

    # Compte total (approximate — pour la pagination)
    count_q = (
        sb.table("contact")
        .select("id", count="exact")
        .eq("tenant_id", tenant_id)
        .is_("deleted_at", "null")
    )
    if q:
        count_q = count_q.or_(
            f"first_name.ilike.%{q}%,last_name.ilike.%{q}%,email.ilike.%{q}%,phone.ilike.%{q}%"
        )
    if category:
        count_q = count_q.eq("category", category)
    if segment:
        count_q = count_q.eq("segment", segment)
    total = (count_q.execute()).count or 0

    return {"contacts": contacts, "total": total, "offset": offset, "limit": limit}


@router.get("/{contact_id}")
async def get_contact(
    contact_id: str,
    tenant_id: str = Depends(get_current_tenant),
):
    """Fiche contact unifiée : infos + tags + leads + RDV + conversations + valeur + reminders."""
    sb = get_supabase_admin()

    res = (
        sb.table("contact")
        .select("*")
        .eq("id", contact_id)
        .eq("tenant_id", tenant_id)
        .is_("deleted_at", "null")
        .single()
        .execute()
    )
    if not res.data:
        raise HTTPException(404, "Contact introuvable")
    contact = res.data

    photo_path = (contact.get("contact_details") or {}).get("photo_path")
    contact["photo_signed_url"] = _signed_url_bucket(sb, PHOTO_BUCKET, photo_path) if photo_path else None

    # Tags
    tag_links = (
        sb.table("contact_tag_link")
        .select("contact_tag(id, name, color)")
        .eq("contact_id", contact_id)
        .execute()
    ).data or []
    contact["tags"] = [l["contact_tag"] for l in tag_links if l.get("contact_tag")]

    # Leads
    leads = (
        sb.table("lead")
        .select("id, status, source, request_type, notes, internal_note, created_at, updated_at")
        .eq("contact_id", contact_id)
        .eq("tenant_id", tenant_id)
        .order("created_at", desc=True)
        .execute()
    ).data or []
    contact["leads"] = leads

    # RDV
    appointments = (
        sb.table("appointment")
        .select("id, status, scheduled_at, end_at, notes, service_offer(name, price_eur)")
        .eq("contact_id", contact_id)
        .order("scheduled_at", desc=True)
        .execute()
    ).data or []
    contact["appointments"] = appointments

    # Valeur totale (somme des prix des RDV confirmés)
    confirmed_prices = [
        float((a.get("service_offer") or {}).get("price_eur") or 0)
        for a in appointments
        if a["status"] == "confirmed"
    ]
    contact["client_value"] = round(sum(confirmed_prices), 2)
    contact["appointments_count"] = len([a for a in appointments if a["status"] == "confirmed"])

    # Inactivité
    confirmed_dates = [
        a["scheduled_at"] for a in appointments if a["status"] == "confirmed"
    ]
    last = max(confirmed_dates) if confirmed_dates else None
    contact["last_appointment_at"] = last
    cutoff = (datetime.now(timezone.utc) - timedelta(days=INACTIVE_DAYS)).isoformat()
    contact["is_inactive"] = (last is None or last < cutoff)

    # Conversations agent IA (5 dernières)
    convs = (
        sb.table("conversation")
        .select("id, agent_type, channel, started_at")
        .eq("contact_id", contact_id)
        .eq("tenant_id", tenant_id)
        .is_("deleted_at", "null")
        .order("started_at", desc=True)
        .limit(5)
        .execute()
    ).data or []
    contact["conversations"] = convs

    # Relances
    reminders = (
        sb.table("contact_reminder")
        .select("id, due_date, note, done, reminder_type, auto_send, sent_at, created_at")
        .eq("contact_id", contact_id)
        .eq("tenant_id", tenant_id)
        .order("due_date")
        .execute()
    ).data or []
    contact["reminders"] = reminders

    return contact


@router.patch("/{contact_id}")
async def update_contact(
    contact_id: str,
    body: ContactUpdate,
    tenant_id: str = Depends(get_current_tenant),
    user: dict = Depends(get_current_user),
):
    """Met à jour les champs éditables d'un contact."""
    sb = get_supabase_admin()

    updates = {k: v for k, v in body.model_dump().items() if v is not None}
    if not updates:
        raise HTTPException(400, "Aucun champ à mettre à jour")

    if "category" in updates and updates["category"] not in CATEGORIES:
        raise HTTPException(400, f"Catégorie invalide (attendu : {', '.join(sorted(CATEGORIES))})")

    if "contact_details" in updates:
        for phone_key in ("phone_pro", "phone_perso"):
            raw = updates["contact_details"].get(phone_key)
            cleaned = clean_phone(raw)
            if raw and not is_valid_phone(cleaned):
                raise HTTPException(400, "Le téléphone ne doit contenir que des chiffres et éventuellement un + au début.")
            if cleaned is not None:
                updates["contact_details"][phone_key] = cleaned

    current_res = (
        sb.table("contact")
        .select("first_name, last_name, email, phone, category, segment, contact_details")
        .eq("id", contact_id)
        .eq("tenant_id", tenant_id)
        .maybe_single()
        .execute()
    )
    current = current_res.data if current_res else None
    if not current:
        raise HTTPException(404, "Contact introuvable")

    if "contact_details" in updates:
        updates["contact_details"] = {**(current.get("contact_details") or {}), **updates["contact_details"]}

    changes = [
        f"{field} : « {current.get(field) or ''} » → « {new_value or ''} »"
        for field, new_value in updates.items()
        if field != "contact_details" and str(current.get(field) or "") != str(new_value or "")
    ]
    if "contact_details" in updates and updates["contact_details"] != (current.get("contact_details") or {}):
        changes.append("informations complémentaires mises à jour")

    res = (
        sb.table("contact")
        .update(updates)
        .eq("id", contact_id)
        .eq("tenant_id", tenant_id)
        .execute()
    )
    if not res.data:
        raise HTTPException(404, "Contact introuvable")

    if changes:
        log_activity(
            tenant_id, user["sub"], "Contact modifié", "; ".join(changes),
            target_type="contact", target_id=contact_id,
        )

    return res.data[0]


@router.get("/{contact_id}/audit-log")
async def get_contact_audit_log(
    contact_id: str,
    limit: int = Query(default=50, le=200),
    tenant_id: str = Depends(get_current_tenant),
):
    """Journal des modifications apportées à la fiche (qui/quoi/quand)."""
    sb = get_supabase_admin()
    rows = (
        sb.table("activity_log")
        .select("id, action, detail, user_id, created_at")
        .eq("tenant_id", tenant_id)
        .eq("target_type", "contact")
        .eq("target_id", contact_id)
        .order("created_at", desc=True)
        .limit(limit)
        .execute()
    ).data or []
    return rows


@router.get("/{contact_id}/activities")
async def list_activities(
    contact_id: str,
    tenant_id: str = Depends(get_current_tenant),
):
    sb = get_supabase_admin()
    rows = (
        sb.table("contact_activity")
        .select("*")
        .eq("contact_id", contact_id)
        .eq("tenant_id", tenant_id)
        .order("created_at", desc=True)
        .execute()
    ).data or []
    return rows


@router.post("/{contact_id}/activities")
async def create_activity(
    contact_id: str,
    body: ActivityCreate,
    tenant_id: str = Depends(get_current_tenant),
):
    if not body.content.strip():
        raise HTTPException(400, "Le contenu ne peut pas être vide")
    if body.type not in ("note", "call", "email", "meeting"):
        raise HTTPException(400, "Type invalide")
    sb = get_supabase_admin()
    row = (
        sb.table("contact_activity")
        .insert({"tenant_id": tenant_id, "contact_id": contact_id,
                 "type": body.type, "content": body.content.strip()})
        .execute()
    ).data[0]
    return row


@router.delete("/{contact_id}/activities/{activity_id}", status_code=204)
async def delete_activity(
    contact_id: str,
    activity_id: str,
    tenant_id: str = Depends(get_current_tenant),
):
    sb = get_supabase_admin()
    sb.table("contact_activity").delete().eq("id", activity_id).eq("tenant_id", tenant_id).execute()


@router.post("/import")
async def import_contacts_csv(
    file: UploadFile = File(...),
    tenant_id: str = Depends(get_current_tenant),
):
    """
    Importe des contacts depuis un fichier CSV ou Excel (.xlsx).
    Colonnes acceptées (insensible à la casse, alias fr tolérés) : voir CSV_COLUMNS.
    Seuls first_name/last_name/email sont obligatoires (au moins un des trois).
    Ignore les doublons (même email déjà présent pour ce tenant).
    """
    filename = (file.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx")):
        raise HTTPException(400, "Le fichier doit être au format CSV (.csv) ou Excel (.xlsx)")

    content = await file.read()

    if filename.endswith(".xlsx"):
        try:
            rows = _read_xlsx_rows(content)
        except Exception as e:
            raise HTTPException(400, f"Fichier Excel illisible : {e}")
    else:
        try:
            text = content.decode("utf-8-sig")  # gère le BOM Excel
        except UnicodeDecodeError:
            text = content.decode("latin-1")

        reader = csv.DictReader(io.StringIO(text))
        # Normalise les noms de colonnes en minuscules sans espaces
        rows = []
        for row in reader:
            rows.append({k.strip().lower(): v.strip() for k, v in row.items() if v is not None})

    if not rows:
        return {"created": 0, "skipped": 0, "errors": []}

    sb = get_supabase_admin()

    # Emails déjà présents pour ce tenant
    existing_emails = {
        r["email"].lower()
        for r in (
            sb.table("contact")
            .select("email")
            .eq("tenant_id", tenant_id)
            .is_("deleted_at", "null")
            .execute()
        ).data or []
        if r.get("email")
    }

    created, skipped = 0, 0
    errors: list[str] = []
    to_insert: list[dict] = []

    for i, row in enumerate(rows, start=2):  # ligne 1 = header
        fields = _csv_row_to_contact_fields(row)

        if not fields["first_name"] and not fields["last_name"] and not fields["email"]:
            skipped += 1
            continue

        if fields["email"] and fields["email"] in existing_emails:
            skipped += 1
            continue

        to_insert.append({
            "tenant_id": tenant_id,
            "source": "csv_import",
            **fields,
        })
        if fields["email"]:
            existing_emails.add(fields["email"])

    if to_insert:
        try:
            sb.table("contact").insert(to_insert).execute()
            created = len(to_insert)
        except Exception as e:
            raise HTTPException(500, f"Erreur lors de l'insertion : {e}")

    return {"created": created, "skipped": skipped, "errors": errors}
